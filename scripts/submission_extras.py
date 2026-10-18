"""Nexus optional initial / required final summary materials, from actual results."""
import json
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml.ns import qn
from topoaml.data import ROOT


def main():
    highlights=['Historical graph rules exclude later-step transaction neighbors.',
                'Static Betti counts repeat conventional structural information.',
                'Persistence does not establish an incremental benefit in the primary test.',
                'Elliptic2 is evaluated as a separate suspicious-subgraph task.']
    assert all(len(s)<=85 for s in highlights)
    blurb=('TopoAML tests persistent graph features under explicit historical observation and structural controls. The primary Bitcoin transaction comparison does not establish an improvement. A separate subgraph experiment preserves its own labels and protocol. The study provides reproducible evidence while exposing limits of coarse timestamps and feature redundancy.')
    assert len(blurb.split())<=50
    doc=Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.add_heading('Highlights',0)
    for item in highlights:doc.add_paragraph(item,style='List Bullet')
    doc.add_heading('eTOC blurb',1);doc.add_paragraph(blurb)
    doc.save(ROOT/'manuscript/highlights_etoc.docx')
    (ROOT/'manuscript/highlights_etoc.md').write_text('# Highlights\n\n'+'\n'.join('- '+h for h in highlights)+'\n\n# eTOC blurb\n\n'+blurb+'\n')
    primary=json.loads((ROOT/'results/primary_comparison.json').read_text())
    plt.rcParams['font.family']='Arial'
    fig=plt.figure(figsize=(4,4),dpi=300);ax=fig.add_axes([0,0,1,1]);ax.axis('off')
    ax.text(.5,.94,'Does topology add information?',ha='center',va='top',fontsize=14,fontweight='bold',color='#19334D')
    ax.text(.5,.82,'Bitcoin transaction graphs',ha='center',fontsize=13)
    ax.text(.5,.73,'Historical neighborhoods\n+ structural controls',ha='center',va='center',fontsize=12,
            bbox=dict(facecolor='#e8eef5',edgecolor='#19334D',boxstyle='round,pad=.4'))
    ax.annotate('',xy=(.5,.57),xytext=(.5,.65),arrowprops=dict(arrowstyle='->',color='#19334D',lw=1.4))
    ax.text(.5,.52,'Add persistent homology',ha='center',fontsize=13)
    ax.text(.5,.40,f"AP change: {primary['delta']:+.5f}",ha='center',fontsize=14,fontweight='bold')
    ax.text(.5,.32,f"95% interval\n[{primary['ci95'][0]:.5f}, {primary['ci95'][1]:.5f}]",ha='center',va='center',fontsize=12)
    ax.text(.5,.19,'No clear incremental benefit',ha='center',fontsize=13,fontweight='bold',color='#884b22')
    ax.text(.5,.07,'Five-seed XGBoost ensembles\nFuture steps 40-49; step-close observation',ha='center',va='center',fontsize=12)
    fig.savefig(ROOT/'manuscript/graphical_abstract.png',dpi=300)
    fig.savefig(ROOT/'manuscript/graphical_abstract.svg')
    plt.close(fig)
    print('Highlights, eTOC blurb and 1200-square graphical abstract written')

if __name__=='__main__':main()
